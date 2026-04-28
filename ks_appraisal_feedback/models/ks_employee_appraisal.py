# -*- coding: utf-8 -*-

from odoo import api, fields, models, Command, _
from odoo.exceptions import UserError, ValidationError
from datetime import date, datetime
import io
import base64

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


class KsEmployeeAppraisal(models.Model):
    _name = 'ks.employee.appraisal'
    _description = 'Employee Appraisal'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'date desc, id desc'
    _rec_name = 'display_name'

    name = fields.Char(
        string='Reference',
        required=True,
        copy=False,
        readonly=True,
        default='New',
        translate=False,
    )
    
    display_name = fields.Char(
        string='Display Name',
        compute='_compute_display_name',
        store=True,
        translate=False,
    )
    
    employee_id = fields.Many2one(
        'hr.employee',
        string='Employee',
        required=True,
        tracking=True,
        index=True,
    )

    employee_code = fields.Char(
        string='Employee Code',
        related='employee_id.employee_code',
        store=True,
        readonly=True,
    )

    manager_id = fields.Many2one(
        'res.users',
        string='Manager',
        required=True,
        tracking=True,
    )

    date = fields.Date(
        string='Appraisal Date',
        required=True,
        default=fields.Date.today,
        tracking=True,
    )

    appraisal_cycle = fields.Selection([
        ('april', 'April'),
        ('october', 'October'),
    ], string='Appraisal Cycle',
        required=True,
        tracking=True,
    )

    job_id = fields.Many2one(
        'hr.job',
        string='Designation',
        related='employee_id.job_id',
        store=True,
        readonly=True,
    )

    department_id = fields.Many2one(
        'hr.department',
        string='Department',
        related='employee_id.department_id',
        store=True,
        readonly=True,
    )

    state = fields.Selection([
        ('draft', 'Draft'),
        ('sent', 'Sent'),
        ('in_progress', 'In Progress'),
        ('completed', 'Completed'),
        ('cancelled', 'Cancelled'),
    ], string='Status',
        default='draft',
        required=True,
        tracking=True,
    )

    feedback_line_ids = fields.One2many(
        'ks.appraisal.feedback.line',
        'appraisal_id',
        string='Feedback Lines',
        copy=True,
    )

    sent_date = fields.Datetime(
        string='Sent Date',
        readonly=True,
    )

    completed_date = fields.Datetime(
        string='Completed Date',
        readonly=True,
    )

    deadline_date = fields.Date(
        string='Deadline',
        tracking=True,
    )

    notes = fields.Text(
        string='Notes',
    )

    question_count = fields.Integer(
        string='Question Count',
        compute='_compute_progress',
    )

    answered_count = fields.Integer(
        string='Answered Count',
        compute='_compute_progress',
    )

    progress = fields.Float(
        string='Progress',
        compute='_compute_progress',
    )

    company_id = fields.Many2one(
        'res.company',
        string='Company',
        default=lambda self: self.env.company,
    )

    @api.depends('feedback_line_ids', 'feedback_line_ids.answer_display', 'feedback_line_ids.display_type')
    def _compute_progress(self):
        for record in self:
            question_lines = record.feedback_line_ids.filtered(lambda l: not l.display_type)
            record.question_count = len(question_lines)
            answered_lines = question_lines.filtered(lambda l: l.answer_display)
            record.answered_count = len(answered_lines)
            if record.question_count:
                record.progress = (record.answered_count / record.question_count) * 100
            else:
                record.progress = 0.0

    @api.depends('name', 'employee_id', 'appraisal_cycle', 'date')
    def _compute_display_name(self):
        for record in self:
            if record.employee_id and record.appraisal_cycle:
                year = record.date.year if record.date else datetime.now().year
                cycle = dict(record._fields['appraisal_cycle'].selection).get(record.appraisal_cycle, '')
                record.display_name = f"{record.employee_id.name} - {cycle} {year}"
            elif record.name and record.name != 'New':
                record.display_name = record.name
            else:
                record.display_name = f"Appraisal #{record.id or 'New'}"

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', 'New') == 'New':
                vals['name'] = self.env['ir.sequence'].next_by_code('ks.employee.appraisal') or 'New'
        return super().create(vals_list)

    @api.onchange('employee_id')
    def _onchange_employee_id(self):
        if self.employee_id:
            if self.employee_id.parent_id and self.employee_id.parent_id.user_id:
                self.manager_id = self.employee_id.parent_id.user_id
            if self.employee_id.appraisal_cycle:
                self.appraisal_cycle = self.employee_id.appraisal_cycle

    def action_populate_questions(self):
        """Populate feedback lines with questions from question master (no section lines)."""
        self.ensure_one()

        if self.state != 'draft':
            raise UserError(_('Questions can only be populated in draft state.'))

        self.feedback_line_ids.unlink()

        categories = self.env['ks.appraisal.question.category'].search([], order='sequence, name')
        lines_to_create = []
        sequence = 10

        for category in categories:
            questions = self.env['ks.question.master'].search([
                ('category_id', '=', category.id),
                ('active', '=', True)
            ], order='sequence, name')
            for question in questions:
                lines_to_create.append(Command.create({
                    'sequence': sequence,
                    'question_id': question.id,
                    'name': question.name,
                }))
                sequence += 10

        if not lines_to_create:
            raise UserError(_('No active questions found. Please configure questions in Question Master first.'))

        self.write({'feedback_line_ids': lines_to_create})

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Questions Populated'),
                'message': _('%d questions added.') % len(lines_to_create),
                'type': 'success',
                'sticky': False,
            }
        }

    def action_send_feedback(self):
        """Send feedback request to manager"""
        self.ensure_one()

        if not self.feedback_line_ids:
            raise UserError(_('Please populate questions first.'))

        if not self.manager_id.email:
            raise UserError(_('Manager does not have an email address.'))

        base_url = self.env['ir.config_parameter'].sudo().get_param('web.base.url').rstrip('/')

        # Get the client action URL for Manager Feedback
        feedback_action = self.env.ref('ks_appraisal_feedback.action_manager_feedback', raise_if_not_found=False)
        if feedback_action:
            feedback_url = f"{base_url}/web#action={feedback_action.id}"
        else:
            feedback_url = f"{base_url}/web#action=ks_appraisal_feedback.action_manager_feedback"

        email_body = f"""
        <div style="font-family: Arial, sans-serif; line-height: 1.6; color: #333;">
            <p>Dear {self.manager_id.name},</p>
            <p>You have been requested to provide appraisal feedback for:</p>

            <table style="margin: 20px 0; border-collapse: collapse;">
                <tr><td style="padding: 8px; font-weight: bold;">Employee:</td><td style="padding: 8px;">{self.employee_id.name}</td></tr>
                <tr><td style="padding: 8px; font-weight: bold;">Employee Code:</td><td style="padding: 8px;">{self.employee_code or 'N/A'}</td></tr>
                <tr><td style="padding: 8px; font-weight: bold;">Designation:</td><td style="padding: 8px;">{self.job_id.name if self.job_id else 'N/A'}</td></tr>
                <tr><td style="padding: 8px; font-weight: bold;">Appraisal Cycle:</td><td style="padding: 8px;">{dict(self._fields['appraisal_cycle'].selection).get(self.appraisal_cycle, '')}</td></tr>
                <tr><td style="padding: 8px; font-weight: bold;">Deadline:</td><td style="padding: 8px;">{self.deadline_date.strftime('%B %d, %Y') if self.deadline_date else 'Not specified'}</td></tr>
            </table>

            <p><a href="{feedback_url}" style="background-color: #875A7B; color: white; padding: 12px 24px; text-decoration: none; border-radius: 4px; display: inline-block; font-weight: bold;">Fill Feedback Form</a></p>

            {f'<p><strong>Notes:</strong> {self.notes}</p>' if self.notes else ''}

            <p>Thank you.</p>
            <p>Best regards,<br/>{self.env.user.name}</p>
        </div>
        """

        try:
            mail = self.env['mail.mail'].create({
                'subject': f'Appraisal Feedback Request: {self.employee_id.name}',
                'body_html': email_body,
                'email_to': self.manager_id.email,
                'email_from': self.env.user.email_formatted or 'noreply@odoo.com',
                'auto_delete': False,
            })
            mail.send()
        except Exception as e:
            raise UserError(_('Failed to send email: %s') % str(e))

        self.write({
            'state': 'sent',
            'sent_date': fields.Datetime.now(),
        })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Request Sent'),
                'message': _('Feedback request sent to %s') % self.manager_id.name,
                'type': 'success',
                'sticky': False,
            }
        }

    def action_resend_feedback(self):
        """Resend feedback request"""
        return self.action_send_feedback()

    def action_start_filling(self):
        """Mark as in progress"""
        self.ensure_one()
        if self.state == 'sent':
            self.write({'state': 'in_progress'})

    def action_mark_completed(self):
        """Mark feedback as completed"""
        self.ensure_one()

        required_lines = self.feedback_line_ids.filtered(
            lambda l: not l.display_type and l.is_required and not l.answer_display
        )

        if required_lines:
            raise UserError(_(
                'Please answer all required questions: %s'
            ) % ', '.join(required_lines.mapped('question_id.name')))

        self.write({
            'state': 'completed',
            'completed_date': fields.Datetime.now(),
        })

    def action_cancel(self):
        """Cancel the appraisal"""
        self.write({'state': 'cancelled'})

    def action_reset_to_draft(self):
        """Reset to draft"""
        self.write({
            'state': 'draft',
            'sent_date': False,
            'completed_date': False,
        })

    @api.model
    def _check_overdue(self):
        """Cron method for overdue appraisals"""
        overdue = self.search([
            ('state', 'in', ['sent', 'in_progress']),
            ('deadline_date', '<', date.today()),
        ])
        if overdue:
            import logging
            _logger = logging.getLogger(__name__)
            _logger.info('Found %d overdue appraisals', len(overdue))

    def action_export_excel(self):
        """Export to Excel with category headers merged above questions"""
        if not self:
            raise UserError(_('Please select at least one record.'))

        if not Workbook:
            raise UserError(_('openpyxl library is required.'))

        wb = Workbook()
        ws = wb.active
        ws.title = "Appraisal Feedback"

        # Styles
        category_fill = PatternFill(start_color="875A7B", end_color="875A7B", fill_type="solid")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        center_alignment = Alignment(horizontal='center', vertical='center')
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Get categories with their questions
        categories = self.env['ks.appraisal.question.category'].search([], order='sequence, name')
        questions_by_category = {}
        all_question_ids = []

        for category in categories:
            questions = self.env['ks.question.master'].search([
                ('category_id', '=', category.id),
                ('active', '=', True)
            ], order='sequence, name')
            if questions:
                questions_by_category[category.id] = {
                    'name': category.name,
                    'questions': questions,
                }
                all_question_ids.extend(questions.ids)

        # Basic headers (columns 1-7)
        basic_headers = ['Employee', 'Employee Code', 'Designation', 'Manager', 'Appraisal Cycle', 'Date', 'Status']
        num_basic_cols = len(basic_headers)

        # Row 1: Category headers (merged cells for questions in each category)
        # Basic columns - merge rows 1 and 2
        for col, header in enumerate(basic_headers, 1):
            # Merge cells for basic headers (row 1 and 2)
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border
            # Also set border for row 2
            ws.cell(row=2, column=col).border = border

        # Category headers - row 1, merged across question columns
        current_col = num_basic_cols + 1
        for cat_id, cat_data in questions_by_category.items():
            num_questions = len(cat_data['questions'])
            if num_questions > 1:
                ws.merge_cells(
                    start_row=1, start_column=current_col,
                    end_row=1, end_column=current_col + num_questions - 1
                )
            cell = ws.cell(row=1, column=current_col, value=cat_data['name'])
            cell.fill = category_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border

            # Set border for all merged cells in category header
            for c in range(current_col, current_col + num_questions):
                ws.cell(row=1, column=c).border = border

            current_col += num_questions

        # Row 2: Question headers
        current_col = num_basic_cols + 1
        for cat_id, cat_data in questions_by_category.items():
            for question in cat_data['questions']:
                cell = ws.cell(row=2, column=current_col, value=question.name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
                current_col += 1

        # Data rows (starting from row 3)
        for row_idx, appraisal in enumerate(self, 3):
            answers = {
                line.question_id.id: line.answer_display
                for line in appraisal.feedback_line_ids
                if line.question_id
            }

            # Basic data
            row_data = [
                appraisal.employee_id.name,
                appraisal.employee_code or '',
                appraisal.job_id.name if appraisal.job_id else '',
                appraisal.manager_id.name,
                dict(appraisal._fields['appraisal_cycle'].selection).get(appraisal.appraisal_cycle, ''),
                str(appraisal.date) if appraisal.date else '',
                dict(appraisal._fields['state'].selection).get(appraisal.state, ''),
            ]

            # Write basic data
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = border

            # Write answers for each question
            current_col = num_basic_cols + 1
            for cat_id, cat_data in questions_by_category.items():
                for question in cat_data['questions']:
                    cell = ws.cell(row=row_idx, column=current_col, value=answers.get(question.id, ''))
                    cell.border = border
                    current_col += 1

        # Adjust column widths
        for col in range(1, num_basic_cols + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15

        # Question columns - slightly wider
        total_cols = num_basic_cols + len(all_question_ids)
        for col in range(num_basic_cols + 1, total_cols + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

        # Set row heights
        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 40

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': 'Appraisal_Feedback_Report.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }